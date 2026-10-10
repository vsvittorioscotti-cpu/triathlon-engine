# -*- coding: utf-8 -*-
import sys, os
sys.path.insert(0,'/tmp/claude-0/-home-claude/9a4d6698-6b4d-546b-bf93-5a7477a0c384/scratchpad')
import forza2 as F
from forza_sess2 import W, componi, _spacchetta, ZONA_FAM, NM_FAM
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.formatting.rule import CellIsRule

AR=lambda **k: Font(name='Arial', size=10, **k)
H=Font(name='Arial', size=10, bold=True, color='FFFFFF'); HF=PatternFill('solid', fgColor='44546A')
BOX=Border(*[Side(style='thin', color='BFBFBF')]*4); TIT=Font(name='Arial', size=13, bold=True)
CEN=Alignment(horizontal='center')

def testa(ws, r, intest, larg):
    for i,(t,w) in enumerate(zip(intest,larg),1):
        c=ws.cell(r,i,t); c.font=H; c.fill=HF; c.border=BOX
        c.alignment=Alignment(horizontal='center', vertical='center', wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width=w
    ws.freeze_panes=ws.cell(r+1,1)

wb=Workbook()
ESERC=sorted(F.ESERCIZI, key=lambda x:(x['famiglia'], F.sottocategoria(x), x['nome']))

# ---------- 1. libreria ----------
ws=wb.active; ws.title='Libreria Esercizi'
ws['A1']='Libreria di Strength & Conditioning per il triathlon'; ws['A1'].font=TIT
ws['A2']=('Fusione di due database forniti da Vittorio il 10/10/2026. NON deriva dal libro di triathlon. '
          'Costo di fatica, rischio DOMS, collocazione e alternative sono DERIVATI da regole esplicite '
          '(vedi Legenda), tarati sugli esempi forniti.')
ws['A2'].font=AR(italic=True)
INT=['Esercizio','Famiglia','Sottocategoria','Distretto','Funzione','Attrezzatura',
     'Nuoto','Bici','Corsa','Serie','Reps base','Reps peak','Modo','Recupero (s)',
     'Costo fatica','Rischio DOMS','Collocazione','Livello','Fasi','Durata (min)','Alternative']
testa(ws,4,INT,[30,17,16,22,28,15,6,6,6,6,10,10,6,10,9,9,30,8,7,10,38])
ISO=lambda e:(e['funzione'].startswith('anti-') or 'tensione' in e['funzione'] or e['rip']>=20)
for r,e in enumerate(ESERC, start=5):
    s=F.scheda(e)
    v=[e['nome'],e['famiglia'],s['sottocat'],e['distretto'],e['funzione'],e['attrezzo'],
       e['N'],e['B'],e['C'],e['serie'],s['reps_base'],s['reps_peak'],'sec' if ISO(e) else 'rip',
       e['rec'],s['fatica'],s['doms'],s['collocazione'],e['livello'],e['fasi']]
    for i,x in enumerate(v,1):
        c=ws.cell(r,i,x); c.font=AR(); c.border=BOX
        if i in (7,8,9,10,13,14,15,16,18,19): c.alignment=CEN
    c=ws.cell(r,20,f'=J{r}*(IF(M{r}="sec",N{r}*0+{e["rip"]},{e["rip"]}*3)+N{r})/60')
    c.font=AR(); c.border=BOX; c.number_format='0.0'; c.alignment=CEN
    c=ws.cell(r,21,s['alternative']); c.font=AR(); c.border=BOX
NR=4+len(ESERC)
for col in ('O','P'):
    ws.conditional_formatting.add(f'{col}5:{col}{NR}',
        CellIsRule(operator='greaterThanOrEqual', formula=['4'], fill=PatternFill('solid', fgColor='FFC7CE')))
    ws.conditional_formatting.add(f'{col}5:{col}{NR}',
        CellIsRule(operator='lessThanOrEqual', formula=['1'], fill=PatternFill('solid', fgColor='C6EFCE')))

# ---------- 2. composizione ----------
ws2=wb.create_sheet('Composizione Workout')
ws2['A1']='Composizione dei workout'; ws2['A1'].font=TIT
ws2['A2']='Serie e ripetizioni qui comandano su quelle di default della libreria. "leggero" abbassa il DOMS dell esercizio.'
ws2['A2'].font=AR(italic=True)
testa(ws2,4,['Codice','Workout','Esercizio','Famiglia','Serie','Ripetizioni','Nota',
             'Modo','Recupero (s)','Minuti','DOMS effettivo'],[10,32,30,17,7,11,10,7,12,9,12])
r=5
for cod,nome,bers,fasi,lista in W:
    for t in lista:
        n,se,ri,q=_spacchetta(t)
        m=f"MATCH(C{r},'Libreria Esercizi'!$A$5:$A${NR},0)"
        for i,x in enumerate([cod,nome,n,F.IDX[n]['famiglia'],se,ri,q or ''],1):
            c=ws2.cell(r,i,x); c.font=AR(); c.border=BOX
            if i in (5,6,7): c.alignment=CEN
        c=ws2.cell(r,8,f"=INDEX('Libreria Esercizi'!$M$5:$M${NR},{m})"); c.font=AR(); c.border=BOX; c.alignment=CEN
        c=ws2.cell(r,9,f"=INDEX('Libreria Esercizi'!$N$5:$N${NR},{m})"); c.font=AR(); c.border=BOX; c.alignment=CEN
        c=ws2.cell(r,10,f'=E{r}*(IF(OR(H{r}="sec",F{r}>=20),F{r},F{r}*3)+I{r})/60')
        c.font=AR(); c.border=BOX; c.number_format='0.0'; c.alignment=CEN
        c=ws2.cell(r,11,f'=IF(G{r}="leggero",MIN(2,INDEX(\'Libreria Esercizi\'!$P$5:$P${NR},{m})),'
                        f'INDEX(\'Libreria Esercizi\'!$P$5:$P${NR},{m}))')
        c.font=AR(); c.border=BOX; c.alignment=CEN
        r+=1
NC=r-1

# ---------- 3. workout ----------
ws3=wb.create_sheet('Workout')
ws3['A1']='I sedici workout standardizzati'; ws3['A1'].font=TIT
ws3['A2']=('La collocazione di un workout e quella del suo esercizio piu esigente. '
           'Zona e carico neuro-muscolare sono STIME NOSTRE: il punteggio basato sulla frequenza '
           'cardiaca sottostima per sua natura la forza pesante e la pliometria.')
ws3['A2'].font=AR(italic=True, color='C00000')
testa(ws3,4,['Codice','Workout','Bersaglio','Esercizi','Minuti','Famiglia dominante','Zona',
             'Carico n-m','Costo fatica','DOMS max','Collocazione','Nuoto','Bici','Corsa','Fasi'],
      [10,30,34,9,8,18,6,10,10,9,30,6,6,6,7])
for r,w in enumerate(W, start=5):
    d=componi(w)
    for i,x in enumerate([d['cod'],d['nome'],d['bersaglio']],1):
        c=ws3.cell(r,i,x); c.font=AR(); c.border=BOX
    for col,f_ in ((4,f"=COUNTIFS('Composizione Workout'!$A$5:$A${NC},$A{r})"),
                   (5,f"=SUMIFS('Composizione Workout'!$J$5:$J${NC},'Composizione Workout'!$A$5:$A${NC},$A{r})"),
                   (10,f"=_xlfn.MAXIFS('Composizione Workout'!$K$5:$K${NC},'Composizione Workout'!$A$5:$A${NC},$A{r})")):
        c=ws3.cell(r,col,f_); c.font=AR(); c.border=BOX; c.number_format='0'; c.alignment=CEN
    for i,x in enumerate([d['famiglia'],d['zona'],d['nm'],d['fatica']],6):
        c=ws3.cell(r,i,x); c.font=AR(); c.border=BOX
        if i>6: c.alignment=CEN
    c=ws3.cell(r,11,f'=IF(J{r}>=4,"almeno 48 h prima del lungo",IF(J{r}=3,"nei giorni gia impegnativi","libera, anche nei giorni facili"))')
    c.font=AR(); c.border=BOX
    for i,x in enumerate([d['N'],d['B'],d['C']],12):
        c=ws3.cell(r,i,x); c.font=AR(); c.border=BOX; c.number_format='0.0'; c.alignment=CEN
    c=ws3.cell(r,15,d['fasi']); c.font=AR(); c.border=BOX; c.alignment=CEN

# ---------- 4. uso per fase ----------
ws4=wb.create_sheet('Uso per Fase')
ws4['A1']='Quante sedute per fase, e quali'; ws4['A1'].font=TIT
ws4['A2']='Una seduta puo essere contemporaneamente forza, core, structural e pliometria: non sono quindici sedute.'
ws4['A2'].font=AR(italic=True)
testa(ws4,4,['Componente','Base','Build','Specifico','Peak'],[30,14,14,14,14])
for r,(k,*v) in enumerate([('Forza pesante','2-3','2','1-2','1'),('Pliometria','1-2','1-2','1','minima'),
        ('Core','3','2-3','2','1-2'),('Structural / prehab','2-3','2-3','2','1-2'),
        ('Bodyweight conditioning','0-2','0-1','occasionale','no'),('Mobilita','3+','3+','3+','3+')], start=5):
    c=ws4.cell(r,1,k); c.font=AR(bold=True); c.border=BOX
    for i,x in enumerate(v,2):
        c=ws4.cell(r,i,x); c.font=AR(); c.border=BOX; c.alignment=CEN
ws4['A13']='Workout disponibili per fase'; ws4['A13'].font=Font(name='Arial', size=11, bold=True)
testa(ws4,14,['Fase','Workout'],[30,120])
for r,(f,nome) in enumerate([('B','Base'),('U','Build'),('S','Specifico'),('P','Peak')], start=15):
    c=ws4.cell(r,1,nome); c.font=AR(bold=True); c.border=BOX
    c=ws4.cell(r,2,', '.join(w[0] for w in W if f in w[3])); c.font=AR(); c.border=BOX

# ---------- 5. legenda ----------
ws5=wb.create_sheet('Legenda')
ws5['A1']='Come e stata costruita questa libreria'; ws5['A1'].font=TIT
righe=[('Provenienza','Due database forniti da Vittorio il 10/10/2026, qui fusi. Il primo ha dato le sei famiglie, il secondo le categorie, i campi di fatica e i template. NON derivano dal libro di triathlon.'),
 ('Esercizi','184 unici. Il secondo database ne ha aggiunti 41 al primo; i nomi gia presenti non sono stati duplicati.'),
 ('Trasferimento','Colonne Nuoto / Bici / Corsa: 0 nessuno, 3 molto specifico. Sul workout e la media pesata sui minuti, e serve da controllo: un workout che dichiara di servire il nuoto e segna 1,2 e mal composto.'),
 ('Costo fatica (1-5)','DERIVATO: Prevenzione 1, Core 2, Forza funzionale 3, Forza massima 4 (5 con bilanciere, trap bar o slitta), Pliometria 2 se stiffness, 3 di norma, 5 se eccentrica alta.'),
 ('Rischio DOMS (1-5)','DERIVATO dal costo di fatica, con due correzioni: eccentrica alta (Nordic, balzi, depth jump) va a 5, eccentrica moderata (heel drop, step-down, affondi camminati) a 3, e il monopodalico sotto carico sale di un punto. Una prescrizione "leggero" lo riabbassa a 2.'),
 ('Collocazione','DERIVATA dal DOMS: 4-5 almeno 48 h prima del lungo, 3 nei giorni gia impegnativi, 1-2 libera anche nei giorni facili. Segue il principio hard days hard, easy days easy, e la regola D9 della dottrina.'),
 ('Alternative','DERIVATE: stessa sottocategoria, profilo di disciplina piu vicino, costo di fatica uguale o inferiore. Servono a sostituire un esercizio senza cambiare il bersaglio della seduta.'),
 ('Reps base / peak','Nel peak le ripetizioni scendono e le serie si riducono: si mantiene la forza, non si costruisce.'),
 ('Zona e carico n-m','STIME NOSTRE. La zona cardiaca fa entrare queste sedute nel punteggio di carico insieme alle tre discipline; il carico neuro-muscolare 0-10 esiste perche quel punteggio sottostima la forza pesante, che costa poco in battiti e molto in fatica.'),
 ('Celle da modificare','Serie, ripetizioni e nota nel foglio Composizione Workout; serie e recupero nella Libreria. Tutto il resto e calcolato.')]
ws5.column_dimensions['A'].width=24; ws5.column_dimensions['B'].width=118
for r,(a,b) in enumerate(righe, start=3):
    c=ws5.cell(r,1,a); c.font=AR(bold=True); c.alignment=Alignment(vertical='top')
    c=ws5.cell(r,2,b); c.font=AR(); c.alignment=Alignment(wrap_text=True, vertical='top')
    ws5.row_dimensions[r].height=42

os.makedirs('/mnt/user-data/outputs', exist_ok=True)
FO='/mnt/user-data/outputs/Preparazione_Atletica_Triathlon.xlsx'
wb.save(FO); print('scritto', FO, '| esercizi', len(ESERC), '| workout', len(W))
