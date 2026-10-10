# -*- coding: utf-8 -*-
import re, sys
sys.path.insert(0,'/tmp/claude-0/-home-claude/9a4d6698-6b4d-546b-bf93-5a7477a0c384/scratchpad')
import dati
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

ACC = {'a':'à','e':'è','i':'ì','o':'ò','u':'ù'}
GIORNI = {"Lunedi":"Lunedì","Martedi":"Martedì","Mercoledi":"Mercoledì","Giovedi":"Giovedì","Venerdi":"Venerdì"}
def fix(t):
    if not isinstance(t,str): return t
    t = re.sub(r"([aeiou])'(?![A-Za-z])", lambda m: ACC[m.group(1)], t)
    t = t.replace("Velocita","Velocità")
    for k,v in GIORNI.items(): t = t.replace(k,v)
    return t

ORD = ["Prima Serie","Seconda Serie","Terza Serie","Quarta Serie","Quinta Serie",
       "Sesta Serie","Settima Serie","Ottava Serie"]

wb = Workbook()

# ---------- Stili ----------
FONT = "Arial"
H_FILL = PatternFill("solid", fgColor="1F3864")
H_FONT = Font(name=FONT, size=10, bold=True, color="FFFFFF")
TITLE  = Font(name=FONT, size=14, bold=True, color="1F3864")
BODY   = Font(name=FONT, size=10)
BOLD   = Font(name=FONT, size=10, bold=True)
THIN   = Side(style="thin", color="BFBFBF")
BRD    = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
ALT    = PatternFill("solid", fgColor="F2F5FA")
FILL_MESO = PatternFill("solid", fgColor="FFF2CC")

def header(ws, row, cols):
    for i,c in enumerate(cols,1):
        cell = ws.cell(row=row, column=i, value=c)
        cell.font, cell.fill, cell.border = H_FONT, H_FILL, BRD
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[row].height = 30

# ============ FOGLIO 1: Dettaglio Serie ============
ws = wb.active
ws.title = "Dettaglio Serie"
cols = ["Mesociclo","Settimana","Nota settimana","Pagina","Giorno","Disciplina","Focus sessione",
        "Obiettivo volume sessione","Volume","Unità","Serie","Descrizione della serie"]
header(ws, 1, cols)
r = 2
for s in dati.SESSIONI:
    for i, testo in enumerate(s["serie"]):
        if s["disc"] == "Pesi":
            etichetta = "Prescrizione" if i < 2 else "Nota"
        elif s["disc"] == "Riposo":
            etichetta = "-"
        else:
            etichetta = ORD[i] if i < len(ORD) else f"Serie {i+1}"
        vals = [fix(s["meso"]), s["sett"], fix(s["nota"]), s["pag"], fix(s["giorno"]), s["disc"],
                fix(s["focus"]), fix(s["vol_txt"]), s["vol_num"], s["unita"], etichetta, fix(testo)]
        for j,v in enumerate(vals,1):
            c = ws.cell(row=r, column=j, value=v)
            c.font, c.border = BODY, BRD
            c.alignment = Alignment(vertical="top", wrap_text=(j==12))
            if j in (2,4,9): c.alignment = Alignment(horizontal="center", vertical="top")
            if r % 2 == 0: c.fill = ALT
        r += 1
widths = [34,10,14,8,11,15,20,24,9,8,14,105]
for i,w in enumerate(widths,1): ws.column_dimensions[get_column_letter(i)].width = w
ws.freeze_panes = "E2"
ws.auto_filter.ref = f"A1:L{r-1}"
LAST_DET = r-1

# ============ FOGLIO 2: Riepilogo Sessioni ============
ws2 = wb.create_sheet("Riepilogo Sessioni")
cols2 = ["Mesociclo","Settimana","Nota settimana","Pagina","Giorno","Disciplina","Focus sessione",
         "Obiettivo volume sessione","Volume","Unità","N. serie","Sintesi allenamento"]
header(ws2, 1, cols2)
r = 2
for s in dati.SESSIONI:
    sintesi = " | ".join(fix(t) for t in s["serie"])
    vals = [fix(s["meso"]), s["sett"], fix(s["nota"]), s["pag"], fix(s["giorno"]), s["disc"],
            fix(s["focus"]), fix(s["vol_txt"]), s["vol_num"], s["unita"], len(s["serie"]), sintesi]
    for j,v in enumerate(vals,1):
        c = ws2.cell(row=r, column=j, value=v)
        c.font, c.border = BODY, BRD
        c.alignment = Alignment(vertical="top", wrap_text=(j==12))
        if j in (2,4,9,11): c.alignment = Alignment(horizontal="center", vertical="top")
        if r % 2 == 0: c.fill = ALT
    r += 1
for i,w in enumerate([34,10,14,8,11,15,20,24,9,8,9,120],1):
    ws2.column_dimensions[get_column_letter(i)].width = w
ws2.freeze_panes = "E2"
ws2.auto_filter.ref = f"A1:L{r-1}"
LAST_SES = r-1

# ============ FOGLIO 3: Volumi Settimanali (formule) ============
ws3 = wb.create_sheet("Volumi Settimanali")
ws3["A1"] = "Volumi settimanali per disciplina"
ws3["A1"].font = TITLE
ws3["A2"] = "Valori calcolati con SUMIFS sul foglio 'Riepilogo Sessioni' (colonna Volume). Fonte dati: piano di allenamento trascritto dal libro, pp. 277-294."
ws3["A2"].font = Font(name=FONT, size=9, italic=True, color="666666")
cols3 = ["Mesociclo","Settimana","Nota","Nuoto (m)","Ciclismo (km)","Corsa (km)","Sedute pesi"]
header(ws3, 4, cols3)
settimane = []
for s in dati.SESSIONI:
    key = (s["meso"], s["sett"], s["nota"])
    if key not in settimane: settimane.append(key)
r = 5
RS = "'Riepilogo Sessioni'"
for meso, sett, nota in settimane:
    ws3.cell(row=r, column=1, value=fix(meso))
    ws3.cell(row=r, column=2, value=sett)
    ws3.cell(row=r, column=3, value=fix(nota))
    ws3.cell(row=r, column=4, value=f'=SUMIFS({RS}!$I$2:$I${LAST_SES},{RS}!$B$2:$B${LAST_SES},$B{r},{RS}!$F$2:$F${LAST_SES},"Nuoto")')
    ws3.cell(row=r, column=5, value=f'=SUMIFS({RS}!$I$2:$I${LAST_SES},{RS}!$B$2:$B${LAST_SES},$B{r},{RS}!$F$2:$F${LAST_SES},"Ciclismo")')
    ws3.cell(row=r, column=6, value=f'=SUMIFS({RS}!$I$2:$I${LAST_SES},{RS}!$B$2:$B${LAST_SES},$B{r},{RS}!$F$2:$F${LAST_SES},"Corsa a piedi")')
    ws3.cell(row=r, column=7, value=f'=COUNTIFS({RS}!$B$2:$B${LAST_SES},$B{r},{RS}!$F$2:$F${LAST_SES},"Pesi")')
    for j in range(1,8):
        c = ws3.cell(row=r, column=j); c.font, c.border = BODY, BRD
        c.alignment = Alignment(horizontal="center" if j>1 else "left", vertical="center")
        if j in (4,5,6): c.number_format = '#,##0.0'
    r += 1
tot = r
ws3.cell(row=tot, column=1, value="TOTALE 6 settimane").font = BOLD
for j,col in zip(range(4,8), "DEFG"):
    c = ws3.cell(row=tot, column=j, value=f'=SUM({col}5:{col}{tot-1})')
    c.font, c.border = BOLD, BRD
    c.number_format = '#,##0.0' if j < 7 else '0'
    c.alignment = Alignment(horizontal="center")
for j in (1,2,3):
    ws3.cell(row=tot, column=j).border = BRD
    ws3.cell(row=tot, column=j).fill = FILL_MESO
for j in range(4,8): ws3.cell(row=tot, column=j).fill = FILL_MESO
for i,w in enumerate([34,11,13,14,15,13,13],1):
    ws3.column_dimensions[get_column_letter(i)].width = w

# ============ FOGLIO 4: Legenda ============
ws4 = wb.create_sheet("Legenda")
ws4["A1"] = "Piano Ironman - Fase di preparazione alla gara (settimane 1-6)"
ws4["A1"].font = TITLE
righe = [
    ("Fonte", "Trascrizione delle pagine 277-294 del libro (sezione 'Distanza Ironman'), fotografate dall'utente."),
    ("Copertura", "Primo Mesociclo: settimane 1, 2, 3 e 4 (Ripristino). Secondo Mesociclo: settimane 5 e 6."),
    ("", ""),
    ("Foglio 'Dettaglio Serie'", "Una riga per ogni serie di ogni sessione, con il testo integrale della descrizione."),
    ("Foglio 'Riepilogo Sessioni'", "Una riga per sessione, con tutte le serie concatenate nella colonna 'Sintesi allenamento'."),
    ("Foglio 'Volumi Settimanali'", "Totali per disciplina calcolati con formule SUMIFS/COUNTIFS sul foglio Riepilogo."),
    ("", ""),
    ("Struttura settimana tipo", "Lunedì: riposo - Martedì: corsa + ciclismo (velocità) - Mercoledì: nuoto (velocità) + pesi (forza) - Giovedì: ciclismo + corsa (andatura di gara) - Venerdì: nuoto (andatura di gara) - Sabato: ciclismo (endurance) + pesi (potenza ed endurance muscolare) - Domenica: nuoto + corsa (endurance)."),
    ("Colonna 'Volume'", "Valore numerico dell'obiettivo di quantità complessiva della sessione: metri per il nuoto, chilometri per ciclismo e corsa. Vuota per le sedute di pesi e per il riposo."),
    ("Ultima serie", "In quasi tutte le sessioni l'ultima serie e' del lavoro aerobico supplementare per raggiungere l'obiettivo di volume della sessione."),
    ("Nota settimana 1", "Nelle sedute con i pesi della settimana 1 il libro specifica: nessuna progressione nei pesi rispetto al periodo di base."),
    ("Nota settimana 4", "Settimana di Ripristino (scarico): volumi e serie ridotti circa del 50%."),
]
r = 3
for k, v in righe:
    a = ws4.cell(row=r, column=1, value=k); a.font = BOLD; a.alignment = Alignment(vertical="top")
    b = ws4.cell(row=r, column=2, value=fix(v)); b.font = BODY
    b.alignment = Alignment(vertical="top", wrap_text=True)
    r += 1
ws4.column_dimensions["A"].width = 30
ws4.column_dimensions["B"].width = 110

out = "/mnt/user-data/outputs/Piano_Ironman_Preparazione_Gara_Sett_1-6.xlsx"
wb.save(out)
print("salvato:", out, "| righe dettaglio:", LAST_DET, "| sessioni:", LAST_SES-1)
