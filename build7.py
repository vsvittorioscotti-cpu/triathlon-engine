# -*- coding: utf-8 -*-
import re, sys
sys.path.insert(0,'/tmp/claude-0/-home-claude/9a4d6698-6b4d-546b-bf93-5a7477a0c384/scratchpad')
import dati7 as dati5
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

ACC = {'a':'à','e':'è','i':'ì','o':'ò','u':'ù'}
GIORNI = {"Lunedi":"Lunedì","Martedi":"Martedì","Mercoledi":"Mercoledì","Giovedi":"Giovedì","Venerdi":"Venerdì"}
def fix(t):
    if not isinstance(t,str): return t
    t = re.sub(r"([aeiou])'(?![A-Za-z])", lambda m: ACC[m.group(1)], t)
    t = t.replace("Velocita","Velocità")
    for k,v in GIORNI.items(): t = t.replace(k,v)
    return t

ORD = ["Prima Serie","Seconda Serie","Terza Serie","Quarta Serie","Quinta Serie",
       "Sesta Serie","Settima Serie","Ottava Serie"]
FONT = "Arial"
H_FONT = Font(name=FONT, size=10, bold=True, color="FFFFFF")
TITLE  = Font(name=FONT, size=14, bold=True, color="1F3864")
BODY   = Font(name=FONT, size=10)
BOLD   = Font(name=FONT, size=10, bold=True)
THIN   = Side(style="thin", color="BFBFBF")
BRD    = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
FILL_TOT = PatternFill("solid", fgColor="FFF2CC")
COLORI = {"Triathlon Olimpico": ("1F6F45","EFF6F1"),
          "Triathlon Sprint":   ("8A1538","FBEFF3"),
          "Half-Ironman":       ("7B3F00","FBF4EC"),
          "Ironman":            ("1F3864","F2F5FA")}
SIGLA = {"Triathlon Olimpico":"Olimpico", "Triathlon Sprint":"Sprint",
         "Half-Ironman":"Half-Ironman", "Ironman":"Ironman"}

wb = Workbook(); wb.remove(wb.active)

def header(ws, row, cols, hexcol):
    fill = PatternFill("solid", fgColor=hexcol)
    for i,c in enumerate(cols,1):
        cell = ws.cell(row=row, column=i, value=c)
        cell.font, cell.fill, cell.border = H_FONT, fill, BRD
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[row].height = 30

def etichetta(s, i):
    if s["disc"] == "Riposo": return "-"
    if s.get("fase") == "Base":
        return "Prescrizione" if s["disc"] != "Pesi" else ("Prescrizione" if i < 2 else "Movimenti")
    if s["disc"] == "Pesi":
        return "Prescrizione" if (i < 2 and not s.get("descr")) else "Nota"
    if s.get("descr"): return "Descrizione"
    return ORD[i] if i < len(ORD) else f"Serie {i+1}"

COLS = ["Piano","Fase","Mesociclo","Settimana","Nota settimana","Pagina","Giorno","Disciplina",
        "Focus sessione","Obiettivo volume sessione","Volume","Unità","Serie","Descrizione della serie"]
W = [16,8,32,10,20,8,11,15,22,24,9,8,14,105]

def costruisci(piano):
    dark, light = COLORI[piano]
    sg = SIGLA[piano]
    sess = [s for s in dati5.TUTTE if s["piano"] == piano]
    ALT = PatternFill("solid", fgColor=light)

    ws = wb.create_sheet(f"Dettaglio Serie ({sg})"[:31])
    header(ws, 1, COLS, dark)
    r = 2
    for s in sess:
        for i, testo in enumerate(s["serie"]):
            vals = [piano, s.get("fase","Gara"), fix(s["meso"]), s["sett"], fix(s["nota"]), s["pag"],
                    fix(s["giorno"]), s["disc"], fix(s["focus"]), fix(s["vol_txt"]), s["vol_num"],
                    s["unita"], etichetta(s,i), fix(testo)]
            for j,v in enumerate(vals,1):
                c = ws.cell(row=r, column=j, value=v)
                c.font, c.border = BODY, BRD
                c.alignment = Alignment(vertical="top", wrap_text=(j==14))
                if j in (4,6,11): c.alignment = Alignment(horizontal="center", vertical="top")
                if r % 2 == 0: c.fill = ALT
            r += 1
    for i,w in enumerate(W,1): ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "G2"; ws.auto_filter.ref = f"A1:N{r-1}"

    ws2 = wb.create_sheet(f"Riepilogo Sess. ({sg})"[:31])
    cols2 = COLS[:12] + ["N. serie","Sintesi allenamento"]
    header(ws2, 1, cols2, dark)
    r = 2
    for s in sess:
        vals = [piano, s.get("fase","Gara"), fix(s["meso"]), s["sett"], fix(s["nota"]), s["pag"],
                fix(s["giorno"]), s["disc"], fix(s["focus"]), fix(s["vol_txt"]), s["vol_num"],
                s["unita"], len(s["serie"]), " | ".join(fix(t) for t in s["serie"])]
        for j,v in enumerate(vals,1):
            c = ws2.cell(row=r, column=j, value=v)
            c.font, c.border = BODY, BRD
            c.alignment = Alignment(vertical="top", wrap_text=(j==14))
            if j in (4,6,11,13): c.alignment = Alignment(horizontal="center", vertical="top")
            if r % 2 == 0: c.fill = ALT
        r += 1
    for i,w in enumerate(W[:12]+[9,120],1): ws2.column_dimensions[get_column_letter(i)].width = w
    ws2.freeze_panes = "G2"; ws2.auto_filter.ref = f"A1:N{r-1}"
    LAST = r-1

    ws3 = wb.create_sheet(f"Volumi Sett. ({sg})"[:31])
    ws3["A1"] = f"Volumi settimanali per disciplina - {piano}"; ws3["A1"].font = TITLE
    ws3["A2"] = f"Calcolati con SUMIFS/COUNTIFS sul foglio '{ws2.title}' (colonna Volume). Fonte: piano trascritto dal libro."
    ws3["A2"].font = Font(name=FONT, size=9, italic=True, color="666666")
    header(ws3, 4, ["Mesociclo","Settimana","Nota","Nuoto (m)","Ciclismo (km)","Corsa (km)","Sedute pesi"], dark)
    chiavi = []
    for s in sess:
        k = (s["meso"], s["sett"], s["nota"])
        if k not in chiavi: chiavi.append(k)
    RS = f"'{ws2.title}'"
    r = 5
    for meso, sett, nota in chiavi:
        ws3.cell(row=r, column=1, value=fix(meso))
        ws3.cell(row=r, column=2, value=sett)
        ws3.cell(row=r, column=3, value=fix(nota))
        for col, disc in ((4,"Nuoto"), (5,"Ciclismo"), (6,"Corsa a piedi")):
            ws3.cell(row=r, column=col,
                     value=f'=SUMIFS({RS}!$K$2:$K${LAST},{RS}!$D$2:$D${LAST},$B{r},{RS}!$H$2:$H${LAST},"{disc}")')
        ws3.cell(row=r, column=7,
                 value=f'=COUNTIFS({RS}!$D$2:$D${LAST},$B{r},{RS}!$H$2:$H${LAST},"Pesi")')
        for j in range(1,8):
            c = ws3.cell(row=r, column=j); c.font, c.border = BODY, BRD
            c.alignment = Alignment(horizontal="center" if j>1 else "left", vertical="center")
            if j in (4,5,6): c.number_format = '#,##0.0'
        r += 1
    tot = r
    ws3.cell(row=tot, column=1, value=f"TOTALE {len(chiavi)} settimane").font = BOLD
    for j, col in zip(range(4,8), "DEFG"):
        c = ws3.cell(row=tot, column=j, value=f'=SUM({col}5:{col}{tot-1})')
        c.font, c.border = BOLD, BRD
        c.number_format = '#,##0.0' if j < 7 else '0'
        c.alignment = Alignment(horizontal="center")
    for j in range(1,8):
        ws3.cell(row=tot, column=j).border = BRD
        ws3.cell(row=tot, column=j).fill = FILL_TOT
    for i,w in enumerate([32,11,20,14,15,13,13],1): ws3.column_dimensions[get_column_letter(i)].width = w

for p in ("Triathlon Olimpico", "Triathlon Sprint", "Half-Ironman", "Ironman"):
    costruisci(p)

# ============ LEGENDA ============
ws4 = wb.create_sheet("Legenda")
ws4["A1"] = "Piani di allenamento triathlon - trascrizione dal libro"; ws4["A1"].font = TITLE
righe = [
    ("Fonte", "Trascrizione delle pagine fotografate del libro, sezioni 'Distanza Sprint', 'Distanza Olimpica', 'Distanza Half-Ironman' e 'Distanza Ironman'."),
    ("I quattro piani", "Sono quattro programmi distinti, con volumi e andature di riferimento diversi: hanno fogli separati e non vanno sommati tra loro. Ordine dei fogli: Olimpico, Sprint, Half-Ironman, Ironman."),
    ("", ""),
    ("OLIMPICO - copertura", "Settimane 1-16, pagine 218-249. Primo e Secondo Mesociclo (Fase di base): sett. 1-8 (la 4a e l'8a sono di Ripristino). Terzo Mesociclo (Fase di preparazione alla gara): sett. 9-12 (la 12a e' di Ripristino). Quarto Mesociclo (Fase di preparazione alla gara): sett. 13-16 (la 16a e' di Tapering e si chiude con la gara). ATTENZIONE: la 12a settimana e' INCOMPLETA - manca pag. 237 (venerdi, sabato e domenica)."),
    ("SPRINT - copertura", "Settimane 1-8 complete, pagine 202-215. Primo Mesociclo (Fase di base): sett. 1-4 (la 4a e' di Ripristino). Secondo Mesociclo (Fase di preparazione alla gara): sett. 5-8 (l'8a e' di Tapering e si chiude con la gara)."),
    ("HALF-IRONMAN - copertura", "Settimane 1-8 complete, pagine 252-274. Primo Mesociclo: sett. 1-4 (la 4a e' di Ripristino). Secondo Mesociclo: sett. 5-8 (la 7a e' tapering solo per la Corsa a piedi, la 8a e' tapering completo)."),
    ("IRONMAN - copertura", "Settimane 1-12 complete, pagine 277-311. Primo Mesociclo: sett. 1-4 (la 4a e' di Ripristino). Secondo Mesociclo: sett. 5-9 (l'8a e' di Ripristino). Terzo Mesociclo: sett. 10-12 (la 12a e' di Tapering)."),
    ("", ""),
    ("Colonna 'Fase'", "Base = fase di base, in cui il libro prescrive solo volume e RPE (nell'Olimpico anche i minuti/metri di endurance sulla forza), non le serie. Gara = fase di preparazione alla gara, con le sessioni suddivise in serie numerate. Olimpico e Sprint hanno entrambe le fasi; per Half-Ironman e Ironman il materiale fotografato copre solo la fase di gara."),
    ("Fogli 'Dettaglio Serie'", "Una riga per ogni serie (o prescrizione) di ogni sessione, con il testo integrale."),
    ("Fogli 'Riepilogo Sess.'", "Una riga per sessione, con tutte le serie concatenate nella colonna 'Sintesi allenamento'."),
    ("Fogli 'Volumi Sett.'", "Totali per disciplina calcolati con formule SUMIFS/COUNTIFS sul relativo foglio di riepilogo."),
    ("", ""),
    ("Struttura settimana tipo (fase di gara)", "Lunedì: riposo - Martedì: corsa + ciclismo (velocità) - Mercoledì: nuoto (velocità) + pesi (forza) - Giovedì: ciclismo + corsa (andatura di gara) - Venerdì: nuoto (andatura di gara) - Sabato: ciclismo (endurance) + pesi (potenza ed endurance muscolare) - Domenica: nuoto + corsa (endurance)."),
    ("Struttura settimana tipo (fase di base)", "Stessa distribuzione dei giorni, ma le sessioni sono classificate come 'distanza intermedia' (martedì, mercoledì, giovedì, venerdì) e 'distanza lunga' (sabato e domenica), con intensità espressa in RPE."),
    ("Colonna 'Volume'", "Valore numerico dell'obiettivo di quantità complessiva della sessione: metri per il nuoto, chilometri per ciclismo e corsa. Vuota per le sedute di pesi, per il riposo e per le sessioni sostituite dalla gara. Unica eccezione: Sprint 8a sett., giovedì corsa, dove il libro esprime l'obiettivo in metri (1.600 m); qui è convertito in 1,6 km per coerenza dei totali."),
    ("Ultima serie", "In quasi tutte le sessioni della fase di gara l'ultima serie e' del lavoro aerobico supplementare per raggiungere l'obiettivo di volume della sessione."),
    ("", ""),
    ("Refusi del libro trascritti fedelmente", "1) Olimpico, sedute di ciclismo sulla velocità: l'Ottava Serie recita 'Della corsa aerobica supplementare' invece di 'Del ciclismo aerobico supplementare'. 2) Nelle settimane di tapering il testo stampa 'vari alattati' dove il senso e' 'vari allunghi': trascritto 'allunghi'."),
    ("Sprint - struttura", "Nello Sprint le sessioni non prevedono serie di endurance sulla forza: le sedute sulla velocità hanno 7 serie invece di 8 e quelle di endurance ne hanno 2. Giovedì il libro suggerisce di combinare assieme ciclismo e corsa (sessione 'brick')."),
    ("Olimpico - sett. 10 e 11", "Le foto delle pagg. 231 e 234 erano poco leggibili: i valori sono stati confermati dalle settimane 14 e 15 (pagg. 243 e 246), che il libro replica identiche."),
    ("Unica pagina mancante", "Pag. 237 - Olimpico, coda della 12a settimana (venerdì, sabato e domenica). Tutto il resto dei quattro piani e' completo."),
    ("Nota 'nessuna progressione'", "Il libro riporta 'nessuna progressione nei pesi rispetto al periodo di base' nella prima settimana della fase di gara di ogni piano (Olimpico 9a sett., Sprint 5a e 6a sett., Half-Ironman 1a sett., Ironman 1a sett.) e nella 6a settimana Half-Ironman."),
]
r = 3
for k, v in righe:
    a = ws4.cell(row=r, column=1, value=k); a.font = BOLD; a.alignment = Alignment(vertical="top")
    b = ws4.cell(row=r, column=2, value=fix(v)); b.font = BODY
    b.alignment = Alignment(vertical="top", wrap_text=True)
    r += 1
ws4.column_dimensions["A"].width = 34
ws4.column_dimensions["B"].width = 118

out = "/mnt/user-data/outputs/Piani_Triathlon_Sprint_Olimpico_Half-Ironman_Ironman.xlsx"
wb.save(out)
print("salvato:", out)
print(wb.sheetnames)
