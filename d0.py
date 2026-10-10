# -*- coding: utf-8 -*-
"""D0: la conclusione sull'intensita' regge al modello a tre domini e allo spostamento delle soglie?"""
from openpyxl import load_workbook
from collections import defaultdict
F='/mnt/user-data/outputs/Piani_Triathlon_Sprint_Olimpico_Half-Ironman_Ironman.xlsx'
wb=load_workbook(F, data_only=True)
SA=[r for r in wb['Struttura Allenamento'].iter_rows(min_row=2, values_only=True)]
PAR={r[0]:r[4] for r in wb['Parametri Atleta'].iter_rows(min_row=6, values_only=True) if r[0]}
SOGLIA={'Corsa a piedi':'corsa_10km','Ciclismo':'bici_20km','Nuoto':'nuoto_1600m'}
PIANI=("Triathlon Sprint","Triathlon Olimpico","Half-Ironman","Ironman")

# IF di ogni blocco = passo di soglia / passo del blocco
blocchi=[]
for r in SA:
    disc, zona, tempo = r[8], r[19], r[23]
    if disc not in SOGLIA or not zona or zona not in PAR or not tempo: continue
    IF = PAR[SOGLIA[disc]] / PAR[zona]
    blocchi.append((r[1], r[0], disc, IF, tempo))
print(f"blocchi con IF calcolabile: {len(blocchi)}\n")

def domini(lt1, lt2):
    d=defaultdict(lambda: defaultdict(float))
    for p,_,_,IF,t in blocchi:
        k = 'LOW' if IF < lt1 else ('MOD' if IF < lt2 else 'HIGH')
        d[p][k]+=t
    return d

print("="*74)
print("1. MODELLO A TRE DOMINI  (LOW <LT1 | MOD LT1-LT2 | HIGH >LT2)")
print("   confine di riferimento: LT1 = IF 0,85   LT2 = IF 1,00\n")
d=domini(0.85,1.00)
print(f"{'piano':22s} {'LOW':>8s} {'MOD':>8s} {'HIGH':>8s}   modello")
for p in PIANI:
    t=sum(d[p].values()); lo,mo,hi=[d[p][k]/t*100 for k in ('LOW','MOD','HIGH')]
    mod = "piramidale" if mo>hi else ("polarizzato" if hi>mo else "-")
    print(f"  {p:20s} {lo:7.1f}% {mo:7.1f}% {hi:7.1f}%   {mod}")

print("\n"+"="*74)
print("2. SENSIBILITA': quanto si muove la conclusione spostando le soglie?\n")
print(f"{'LT1':>5s} {'LT2':>5s} | " + " | ".join(f"{p.replace('Triathlon ','')[:9]:>9s}" for p in PIANI))
print(f"{'':>5s} {'':>5s} | " + " | ".join(f"{'LOW/MOD/HI':>9s}" for p in PIANI))
righe=[]
for lt1 in (0.78,0.82,0.85,0.88,0.90):
    for lt2 in (0.96,1.00,1.04):
        d=domini(lt1,lt2); out=[]
        for p in PIANI:
            t=sum(d[p].values()); lo,mo,hi=[d[p][k]/t*100 for k in ('LOW','MOD','HIGH')]
            out.append(f"{lo:.0f}/{mo:.0f}/{hi:.0f}")
            righe.append((lt1,lt2,p,lo,mo,hi))
        print(f"{lt1:5.2f} {lt2:5.2f} | " + " | ".join(f"{o:>9s}" for o in out))

print("\n  -> escursione di LOW su tutte le combinazioni di soglia:")
for p in PIANI:
    v=[r[3] for r in righe if r[2]==p]
    print(f"     {p:22s} da {min(v):5.1f}% a {max(v):5.1f}%   (ampiezza {max(v)-min(v):4.1f} punti)")
print("\n  -> ordinamento dei piani per quota LOW, in ognuna delle 15 combinazioni:")
ord_=set()
for lt1 in (0.78,0.82,0.85,0.88,0.90):
    for lt2 in (0.96,1.00,1.04):
        d=domini(lt1,lt2)
        o=tuple(sorted(PIANI, key=lambda p:-d[p]['LOW']/sum(d[p].values())))
        ord_.add(o)
print(f"     ordinamenti distinti trovati: {len(ord_)}")
for o in ord_: print("     " + "  >  ".join(x.replace('Triathlon ','') for x in o))

print("\n"+"="*74)
print("3. TIME-IN-ZONE contro SESSION-GOAL  (il problema che segnala Seiler)\n")
# session goal: la zona caratterizzante della seduta (blocco piu' intenso non di recupero)
RANK={'Z5':1,'Z4':2,'Z3':3,'Z2':4,'Z1':5}
DOM={'Z5':'HIGH','Z4':'HIGH','Z3':'MOD','Z2':'LOW','Z1':'LOW'}
sess=defaultdict(list); pia={}
for r in SA:
    if r[8] in ('Pesi','Riposo') or not r[30]: continue
    pia[r[0]]=r[1]
    if r[13] not in ('Recupero','Recupero attivo','Aerobico supplementare'):
        sess[r[0]].append(r[30])
sg=defaultdict(lambda: defaultdict(int))
for sid,zs in sess.items():
    if not zs: continue
    sg[pia[sid]][DOM[min(zs,key=lambda z:RANK[z])]]+=1
d=domini(0.85,1.00)
print(f"{'piano':22s} {'time-in-zone LOW/MOD/HIGH':>28s}    {'session-goal LOW/MOD/HIGH':>28s}")
for p in PIANI:
    t=sum(d[p].values()); tz=[d[p][k]/t*100 for k in ('LOW','MOD','HIGH')]
    s=sum(sg[p].values()); sgp=[sg[p][k]/s*100 for k in ('LOW','MOD','HIGH')]
    print(f"  {p:20s} {tz[0]:8.1f}% {tz[1]:7.1f}% {tz[2]:7.1f}%    {sgp[0]:8.1f}% {sgp[1]:7.1f}% {sgp[2]:7.1f}%")
