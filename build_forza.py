# -*- coding: utf-8 -*-
import sys
sys.path.insert(0,'/tmp/claude-0/-home-claude/9a4d6698-6b4d-546b-bf93-5a7477a0c384/scratchpad')
from forza_lib import ESERCIZI, IDX
from forza_sess import SESSIONI, ZONA_FAM, NM_FAM
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

AR  = lambda **k: Font(name='Arial', size=10, **k)
H   = Font(name='Arial', size=10, bold=True, color='FFFFFF')
HF  = PatternFill('solid', fgColor='44546A')
GIA = PatternFill('solid', fgColor='FFF2CC')
BOX = Border(*[Side(style='thin', color='BFBFBF')]*4)
TIT = Font(name='Arial', size=13, bold=True)

def testa(ws, riga, intest, larg):
    for i,(t,w) in enumerate(zip(intest,larg),1):
        c = ws.cell(riga,i,t); c.font=H; c.fill=HF; c.border=BOX
        c.alignment=Alignment(horizontal='center', vertical='center', wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width=w
    ws.freeze_panes = ws.cell(riga+1,1)

wb = Workbook()

# ---------------- 1. libreria ----------------
ws = wb.active; ws.title = 'Libreria Esercizi'
ws['A1']='Libreria di preparazione atletica per il triatleta'; ws['A1'].font=TIT
ws['A2']=('Fonte: impostazione, famiglie e liste fornite da Vittorio (10/10/2026). '
          'NON deriva dal libro di triathlon: e\' materiale indipendente e va trattato come tale.')
ws['A2'].font=AR(italic=True)
ws['A3']=('Trasferimento: 0 = nessuno, 1 = basso, 2 = medio, 3 = alto.  '
          'Fasi: B = base, U = build, S = specifico, P = peak.  '
          'Modo "sec" = tenuta isometrica (le ripetizioni sono secondi), "rip" = ripetizioni.')
ws['A3'].font=AR(italic=True)
INT=['Esercizio','Famiglia','Distretto','Funzione','Attrezzatura','Nuoto','Bici','Corsa',
     'Serie','Ripetizioni','Modo','Recupero (s)','Livello','Fasi','Durata (min)']
testa(ws,5,INT,[30,18,24,30,16,7,7,7,7,11,7,11,9,8,11])
ISO = lambda e: (e['funzione'].startswith('anti-') or 'tensione' in e['funzione'] or e['rip']>=20)
for r,e in enumerate(sorted(ESERCIZI,key=lambda x:(x['famiglia'],x['nome'])), start=6):
    v=[e['nome'],e['famiglia'],e['distretto'],e['funzione'],e['attrezzo'],e['N'],e['B'],e['C'],
       e['serie'],e['rip'],'sec' if ISO(e) else 'rip',e['rec'],e['livello'],e['fasi']]
    for i,x in enumerate(v,1):
        c=ws.cell(r,i,x); c.font=AR(); c.border=BOX
        if i in (6,7,8,9,10,12): c.alignment=Alignment(horizontal='center')
    c=ws.cell(r,15,f'=I{r}*(IF(K{r}="sec",J{r},J{r}*3)+L{r})/60'); c.font=AR(); c.border=BOX
    c.number_format='0.0'; c.alignment=Alignment(horizontal='center')
NR = 5+len(ESERCIZI)

# ---------------- 2. composizione ----------------
ws2 = wb.create_sheet('Composizione Sessioni')
ws2['A1']='Composizione delle sessioni'; ws2['A1'].font=TIT
ws2['A2']='Le serie qui possono differire da quelle di default della libreria: comandano queste.'
ws2['A2'].font=AR(italic=True)
testa(ws2,4,['Codice','Sessione','Esercizio','Famiglia','Serie','Ripetizioni','Modo','Recupero (s)','Minuti'],
      [9,30,30,18,7,11,7,12,9])
r=5
for cod,nome,bers,fasi,lista in SESSIONI:
    for es,serie in lista:
        f=f'MATCH(C{r},\'Libreria Esercizi\'!$A$6:$A${NR},0)'
        for i,x in enumerate([cod,nome,es,IDX[es]['famiglia'],serie],1):
            c=ws2.cell(r,i,x); c.font=AR(); c.border=BOX
        for col,lettera in ((6,'J'),(7,'K'),(8,'L')):
            c=ws2.cell(r,col,f'=INDEX(\'Libreria Esercizi\'!${lettera}$6:${lettera}${NR},{f})')
            c.font=AR(); c.border=BOX; c.alignment=Alignment(horizontal='center')
        c=ws2.cell(r,9,f'=E{r}*(IF(G{r}="sec",F{r},F{r}*3)+H{r})/60'); c.font=AR(); c.border=BOX
        c.number_format='0.0'; c.alignment=Alignment(horizontal='center')
        r+=1
NC=r-1

# ---------------- 3. sessioni ----------------
ws3 = wb.create_sheet('Sessioni')
ws3['A1']='Sessioni di preparazione atletica'; ws3['A1'].font=TIT
ws3['A2']=('Una sessione ha un bersaglio dichiarato, non e\' "palestra". '
           'Durata e numero di esercizi sono calcolati dalla composizione.')
ws3['A2'].font=AR(italic=True)
ws3['A3']=('Zona e carico neuromuscolare sono STIME NOSTRE: il libro non fornisce ne\' FC ne\' RPE '
           'per il lavoro a secco, e il punteggio basato sulla frequenza cardiaca sottostima la forza pesante.')
ws3['A3'].font=AR(italic=True, color='C00000')
testa(ws3,5,['Codice','Sessione','Bersaglio','Esercizi','Minuti','Famiglia dominante','Zona',
             'Carico neuro-muscolare','Nuoto','Bici','Corsa','Fasi'],
      [9,28,34,9,9,20,7,13,7,7,7,8])
for r,(cod,nome,bers,fasi,lista) in enumerate(SESSIONI, start=6):
    fam={}
    for es,n in lista: fam[IDX[es]['famiglia']]=fam.get(IDX[es]['famiglia'],0)+n
    dom=max(fam,key=fam.get)
    for i,x in enumerate([cod,nome,bers],1):
        c=ws3.cell(r,i,x); c.font=AR(); c.border=BOX
    for col,form in ((4,f'=COUNTIFS(\'Composizione Sessioni\'!$A$5:$A${NC},$A{r})'),
                     (5,f'=SUMIFS(\'Composizione Sessioni\'!$I$5:$I${NC},\'Composizione Sessioni\'!$A$5:$A${NC},$A{r})')):
        c=ws3.cell(r,col,form); c.font=AR(); c.border=BOX
        c.number_format='0'; c.alignment=Alignment(horizontal='center')
    for i,x in enumerate([dom,ZONA_FAM[dom],NM_FAM[dom]],6):
        c=ws3.cell(r,i,x); c.font=AR(); c.border=BOX
        if i>6: c.alignment=Alignment(horizontal='center')
    for col,lettera in ((9,'F'),(10,'G'),(11,'H')):
        c=ws3.cell(r,col,
          f'=SUMPRODUCT(--(\'Composizione Sessioni\'!$A$5:$A${NC}=$A{r}),\'Composizione Sessioni\'!$I$5:$I${NC},'
          f'INDEX(\'Libreria Esercizi\'!${lettera}$6:${lettera}${NR},'
          f'MATCH(\'Composizione Sessioni\'!$C$5:$C${NC},\'Libreria Esercizi\'!$A$6:$A${NR},0)))/$E{r}')
        c.font=AR(); c.border=BOX; c.number_format='0.0'; c.alignment=Alignment(horizontal='center')
    c=ws3.cell(r,12,fasi); c.font=AR(); c.border=BOX; c.alignment=Alignment(horizontal='center')

# ---------------- 4. legenda ----------------
ws4 = wb.create_sheet('Legenda')
ws4['A1']='Come si legge e come si usa'; ws4['A1'].font=TIT
righe=[('Provenienza','Questa libreria NON viene dal libro di triathlon. Famiglie, esercizi e impostazione sono stati forniti da Vittorio il 10/10/2026.'),
 ('Le sei famiglie','Forza massima, Forza funzionale, Core e stabilita\', Prevenzione, Pliometria, Bodyweight.'),
 ('Trasferimento','Le colonne Nuoto / Bici / Corsa dicono quanto l\'esercizio serve a quella disciplina: 0 nessuno, 3 alto. Sulle sessioni e\' la media pesata sui minuti.'),
 ('Fasi','B base, U build, S specifico, P peak. Il taper usa solo FT-01.'),
 ('Frequenza','Base 2-3 sessioni a settimana, Build 2, Specifico 1-2, Peak 1 da 20-40 minuti, Taper solo attivazione.'),
 ('Zona cardiaca','STIMA NOSTRA, assegnata per famiglia dominante. Serve a far entrare queste sedute nel punteggio di carico insieme alle altre tre discipline.'),
 ('Carico neuro-muscolare','STIMA NOSTRA, 0-10. Esiste perche\' il punteggio basato sulla frequenza cardiaca sottostima per sua natura la forza pesante e la pliometria: una seduta puo\' costare poco in FC e molto in fatica.'),
 ('Celle da modificare','Le serie nel foglio Composizione Sessioni, e ripetizioni/recupero nella Libreria. Il resto e\' calcolato.'),
 ('Durata','serie x (lavoro + recupero). Il lavoro e\' le ripetizioni x 3 secondi, oppure i secondi stessi per le tenute isometriche (modo "sec").')]
ws4.column_dimensions['A'].width=26; ws4.column_dimensions['B'].width=110
for r,(a,b) in enumerate(righe, start=3):
    c=ws4.cell(r,1,a); c.font=AR(bold=True); c.alignment=Alignment(vertical='top')
    c=ws4.cell(r,2,b); c.font=AR(); c.alignment=Alignment(wrap_text=True, vertical='top')
    ws4.row_dimensions[r].height=32

import os
os.makedirs('/mnt/user-data/outputs', exist_ok=True)
F='/mnt/user-data/outputs/Preparazione_Atletica_Triathlon.xlsx'
wb.save(F); print('scritto', F)
