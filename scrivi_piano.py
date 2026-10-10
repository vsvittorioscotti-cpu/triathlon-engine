# -*- coding: utf-8 -*-
import json, math
from collections import defaultdict
from openpyxl import load_workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.formatting.rule import CellIsRule
d=json.load(open('piano.json')); PIANO=d["PIANO"]; W=d["W"]; A=d["A"]; TAPER_DA=d["TAPER_DA"]
F='/mnt/user-data/outputs/Piani_Triathlon_Sprint_Olimpico_Half-Ironman_Ironman.xlsx'
wb=load_workbook(F)
for _n in list(wb.sheetnames):
    if _n.startswith("Piano Generato") or _n.startswith("Piano - Settimane") or _n.startswith("Validazione Piano"):
        del wb[_n]
FONT="Arial"; BODY=Font(name=FONT,size=10); BOLD=Font(name=FONT,size=10,bold=True)
TITLE=Font(name=FONT,size=14,bold=True,color="1F3864")
TH=Side(style="thin",color="BFBFBF"); BRD=Border(left=TH,right=TH,top=TH,bottom=TH)
HF=Font(name=FONT,size=10,bold=True,color="FFFFFF"); ALT=PatternFill("solid",fgColor="F7F7F7")
def hdr(ws,row,cols,col="404040"):
    f=PatternFill("solid",fgColor=col)
    for i,c in enumerate(cols,1):
        x=ws.cell(row=row,column=i,value=c); x.font=HF; x.fill=f; x.border=BRD
        x.alignment=Alignment(horizontal="center",vertical="center",wrap_text=True)
    ws.row_dimensions[row].height=30

GG=["Lunedì","Martedì","Mercoledì","Giovedì","Venerdì","Sabato","Domenica"]
for p in PIANO: p['giorno_n']=GG.index(p['giorno'])+1; p['day']=(p['w']-1)*7+p['giorno_n']

# ---- foglio Piano Generato ----
ws=wb.create_sheet("Piano Generato")
ws["A1"]=f"Piano generato - {A['gara']}, {A['settimane']} settimane, tetto {A['ore_picco']} h/settimana"; ws["A1"].font=TITLE
ws["A2"]=("Prodotto dal generatore a partire dalla dottrina e dalla libreria di allenamenti. "
          "Ogni seduta e' una variante pescata dal catalogo, scalata sul tempo bersaglio della settimana.")
ws["A2"].font=Font(name=FONT,size=9,italic=True,color="666666")
C=["Settimana","Macrociclo","Blocco","Tipo settimana","Giorno","Disciplina","Stimolo target",
   "Allenamento (codice)","Archetipo base","Durata (min)","Punti Edwards","Punti andatura","Dominio"]
hdr(ws,4,C); r=5
for p in sorted(PIANO,key=lambda x:(x['w'],x['giorno_n'],x['disc'])):
    for j,v in enumerate([p['w'],p['macro'],p['blocco'],p['tipo'],p['giorno'],p['disc'],p['ob'],
                          p['cod'],p['struttura'],round(p['tempo']),round(p['edw']),round(p['rtss']),p['dom']],1):
        c=ws.cell(row=r,column=j,value=v); c.font=BODY; c.border=BRD
        c.alignment=Alignment(horizontal="center" if j in(1,3,10,11,12,13) else "left",vertical="center")
        if r%2==0: c.fill=ALT
    r+=1
PLAST=r-1
for i,w_ in enumerate([11,14,8,14,11,15,24,22,16,12,13,13,10],1): ws.column_dimensions[get_column_letter(i)].width=w_
ws.freeze_panes="E5"; ws.auto_filter.ref=f"A4:M{PLAST}"

# ---- foglio Piano - Settimane con CTL/ATL/TSB ----
gg_tss=defaultdict(float)
for p in PIANO: gg_tss[p['day']]+=p['rtss']
ND=A['settimane']*7; ctl=atl=0.0; serie={}
for dday in range(1,ND+1):
    t=gg_tss.get(dday,0.0)
    ctl+= (t-ctl)/42; atl+= (t-atl)/7
    serie[dday]=(ctl,atl,ctl-atl)
ws2=wb.create_sheet("Piano - Settimane")
ws2["A1"]="Riepilogo settimanale e curve di carico"; ws2["A1"].font=TITLE
ws2["A2"]=("CTL = media esponenziale a 42 giorni dei punti andatura (condizione). ATL = a 7 giorni (fatica). "
           "TSB = CTL - ATL (freschezza). Alla gara il TSB deve essere positivo.")
ws2["A2"].font=Font(name=FONT,size=9,italic=True,color="666666")
hdr(ws2,4,["Settimana","Macrociclo","Blocco","Tipo","Sedute","Ore","Punti Edwards","Punti andatura",
           "% LOW","% MOD","% HIGH","Ore nuoto","Ore bici","Ore corsa","CTL","ATL","TSB"])
sett=defaultdict(lambda: defaultdict(float)); ns=defaultdict(int)
for p in PIANO:
    s=sett[p['w']]; s['ore']+=p['tempo']/60; s['edw']+=p['edw']; s['rtss']+=p['rtss']; s[p['dom']]+=p['tempo']
    s[p['disc']]+=p['tempo']/60; ns[p['w']]+=1
r=5
for x in W:
    s=sett[x['w']]; t=s['LOW']+s['MOD']+s['HIGH'] or 1
    c_,a_,b_=serie[x['w']*7]
    for j,v in enumerate([x['w'],x['macro'],x['blocco'],x['tipo'],ns[x['w']],round(s['ore'],1),
                          round(s['edw']),round(s['rtss']),s['LOW']/t,s['MOD']/t,s['HIGH']/t,
                          round(s['Nuoto'],1),round(s['Ciclismo'],1),round(s['Corsa a piedi'],1),
                          round(c_,1),round(a_,1),round(b_,1)],1):
        c=ws2.cell(row=r,column=j,value=v); c.font=BODY; c.border=BRD
        c.alignment=Alignment(horizontal="center" if j!=2 else "left",vertical="center")
        if j in (9,10,11): c.number_format='0%'
        if x['tipo']=='Scarico': c.fill=PatternFill("solid",fgColor="EDF3FA")
        if x['tipo']=='Taper': c.fill=PatternFill("solid",fgColor="FFF2CC")
    r+=1
WLAST=r-1
for i,w_ in enumerate([11,14,8,11,9,8,14,14,9,9,9,11,11,11,9,9,9],1): ws2.column_dimensions[get_column_letter(i)].width=w_
ws2.freeze_panes="C5"
ws2.conditional_formatting.add(f"Q5:Q{WLAST}",CellIsRule(operator="lessThan",formula=["-10"],
    fill=PatternFill("solid",fgColor="FFC7CE")))
ws2.conditional_formatting.add(f"Q5:Q{WLAST}",CellIsRule(operator="greaterThan",formula=["0"],
    fill=PatternFill("solid",fgColor="C6EFCE")))

# ---- foglio Validazione ----
D=defaultdict(float)
for p in PIANO: D[p['dom']]+=p['tempo']
T=sum(D.values()); low,mod,hi=[D[k]/T*100 for k in ('LOW','MOD','HIGH')]
sc=[x for x in W if x['tipo']=='Scarico']; prof=[]
for x in sc:
    pre=[y for y in W if y['w']<x['w'] and y['tipo']=='Carico']
    if pre: prof.append((sett[x['w']]['ore']/sett[pre[-1]['w']]['ore']-1)*100)
picco=max(sett[w_]['ore'] for w_ in sett)
brick=len({p['w'] for p in PIANO if p['cod']=='BRICK'})
lungo=max((p['tempo'] for p in PIANO if p['giorno']=='Sabato' and p['disc']=='Ciclismo'),default=0)
tapw=[x for x in W if x['tipo']=='Taper']
tsb_gara=serie[ND][2]
REG=[("I4","INVARIANT","Prevalenza di lavoro a bassa intensita'",f"LOW {low:.0f}% (soglia 60%)",low>=60),
     ("R11","DEFAULT","Quota sopra soglia sotto controllo",f"HIGH {hi:.0f}% (tetto 15%)",hi<=15),
     ("I4b","INVARIANT","Forma piramidale, non polarizzata",f"MOD {mod:.0f}% maggiore di HIGH {hi:.0f}%",mod>hi),
     ("D1","DEFAULT","Ciclo carico:scarico 3:1",f"{len(sc)} scarichi, tutti in quarta posizione",all(x['pos']==4 for x in sc)),
     ("D2","DEFAULT","Profondita' dello scarico nel range",f"da {min(prof):.0f}% a {max(prof):.0f}%",all(-55<=p<=-25 for p in prof)),
     ("A3","ADAPTIVE","Tetto orario dell'atleta rispettato",f"picco {picco:.1f} h contro {A['ore_picco']} disponibili",picco<=A['ore_picco']+0.05),
     ("I6","INVARIANT","Transizione bici-corsa allenata",f"{brick} settimane con brick",brick>0),
     ("H2","HEURISTIC","Specificita' del lungo in bici",f"lungo massimo {lungo:.0f} min",lungo>=168),
     ("I5","INVARIANT","Taper presente e progressivo",f"{len(tapw)} settimane, ultima {sett[tapw[-1]['w']]['ore']:.1f} h",len(tapw)>=2),
     ("—","CONTROLLO","Freschezza alla gara (TSB positivo)",f"TSB finale {tsb_gara:+.1f}",tsb_gara>0)]
ws3=wb.create_sheet("Validazione Piano")
ws3["A1"]="Validazione del piano generato contro le regole di costruzione"; ws3["A1"].font=TITLE
ws3["A2"]="Ogni regola del foglio 'Regole di Costruzione' che sia verificabile su un piano viene controllata qui."
ws3["A2"].font=Font(name=FONT,size=9,italic=True,color="666666")
hdr(ws3,4,["ID","Status","Regola","Misura sul piano","Esito"])
r=5
for id_,st,reg,mis,ok in REG:
    for j,v in enumerate([id_,st,reg,mis,"OK" if ok else "VIOLATA"],1):
        c=ws3.cell(row=r,column=j,value=v); c.font=BOLD if j==5 else BODY; c.border=BRD
        c.alignment=Alignment(horizontal="center" if j in(1,2,5) else "left",vertical="center",wrap_text=(j==3))
    r+=1
ws3.cell(row=r+1,column=3,value=f"Regole rispettate: {sum(1 for *_ ,o in REG if o)} su {len(REG)}").font=BOLD
for i,w_ in enumerate([8,14,46,44,12],1): ws3.column_dimensions[get_column_letter(i)].width=w_
ws3.conditional_formatting.add(f"E5:E{r-1}",CellIsRule(operator="equal",formula=['"OK"'],fill=PatternFill("solid",fgColor="C6EFCE")))
ws3.conditional_formatting.add(f"E5:E{r-1}",CellIsRule(operator="equal",formula=['"VIOLATA"'],fill=PatternFill("solid",fgColor="FFC7CE")))

ordine=[n for n in wb.sheetnames if n not in ("Piano Generato","Piano - Settimane","Validazione Piano","Legenda")]
ordine=["Piano Generato","Piano - Settimane","Validazione Piano"]+ordine+["Legenda"]
wb._sheets=[wb[n] for n in ordine]
wb.save(F)
print(f"scritto. sedute {PLAST-4}, settimane {WLAST-4}, regole {len(REG)} ({sum(1 for *_,o in REG if o)} OK)")
print(f"CTL finale {serie[ND][0]:.1f} | ATL {serie[ND][1]:.1f} | TSB {tsb_gara:+.1f}")
print("CTL al picco:", f"{max(serie[dd][0] for dd in serie):.1f}")
