# -*- coding: utf-8 -*-
from openpyxl import load_workbook
from collections import defaultdict, Counter
import statistics as st
F='/mnt/user-data/outputs/Piani_Triathlon_Sprint_Olimpico_Half-Ironman_Ironman.xlsx'
wb=load_workbook(F, data_only=True)
SA=[r for r in wb['Struttura Allenamento'].iter_rows(min_row=2, values_only=True)]
CS=[r for r in wb['Controllo Sessioni'].iter_rows(min_row=5, values_only=True) if r[0]]
OC=[r for r in wb['Obiettivi di Carico'].iter_rows(min_row=6, max_row=49, values_only=True) if r[0] and isinstance(r[2],int)]
PIANI=("Triathlon Sprint","Triathlon Olimpico","Half-Ironman","Ironman")
GIORNI=["Lunedì","Martedì","Mercoledì","Giovedì","Venerdì","Sabato","Domenica"]

print("="*78); print("A. STRUTTURA DEI GIORNI")
sess=defaultdict(dict)
for r in SA:
    sess[r[0]] = dict(piano=r[1],sett=r[4],gg=r[7],disc=r[8],focus=r[9],ob=r[10])
for p in PIANI:
    ss=[v for v in sess.values() if v['piano']==p]
    sett=sorted({v['sett'] for v in ss})
    per_sett=Counter((v['sett']) for v in ss if v['disc']!='Riposo')
    print(f"\n  {p}: {len(sett)} settimane")
    print(f"    sedute/settimana: {sorted(set(per_sett.values()))}  (media {st.mean(per_sett.values()):.1f})")
    gg=defaultdict(Counter)
    for v in ss:
        if v['disc']!='Riposo': gg[v['gg']][v['disc']]+=1
    for g in GIORNI:
        c=gg.get(g)
        if not c: print(f"    {g:10s} RIPOSO"); continue
        tot=sum(c.values()); n=len(sett)
        print(f"    {g:10s} " + ", ".join(f"{k} x{v//n}" for k,v in c.most_common()))

print("\n"+"="*78); print("B. CICLO CARICO / SCARICO")
for p in PIANI:
    w=[r for r in OC if r[0]==p]
    print(f"\n  {p}")
    base=None
    for r in w:
        ore=r[4]; edw=r[5]
        seg=""
        if base: seg=f"  ({(ore/base-1)*100:+5.0f}% vs sett.prec.)"
        print(f"    sett.{r[2]:2d} {str(r[3] or ''):22s} {ore:5.1f} h  {edw:6.0f} pt{seg}")
        base=ore
    ore=[r[4] for r in w]
    scar=[i for i,r in enumerate(w) if r[3] and ('Ripristino' in str(r[3]) or 'Tapering' in str(r[3]))]
    print(f"    -> settimane di scarico in posizione: {[w[i][2] for i in scar]}")
    for i in scar:
        if i>0: print(f"       scarico sett.{w[i][2]}: {(w[i][4]/w[i-1][4]-1)*100:+.0f}% ore, {(w[i][5]/w[i-1][5]-1)*100:+.0f}% carico")

print("\n"+"="*78); print("C. DISTRIBUZIONE PER ZONA (% del tempo)")
z=defaultdict(lambda: defaultdict(float))
for r in SA:
    if r[30] and r[23]: z[r[1]][r[30]]+=r[23]
for p in PIANI:
    tot=sum(z[p].values())
    print(f"\n  {p}  (tot {tot/60:.0f} h)")
    for k in ("Z1","Z2","Z3","Z4","Z5"):
        v=z[p].get(k,0); print(f"    {k}: {v/tot*100:5.1f}%  ({v/60:6.1f} h)")

print("\n"+"="*78); print("D. FOCUS PER MESOCICLO - la disciplina prioritaria cambia?")
for p in PIANI:
    print(f"\n  {p}")
    mes=defaultdict(lambda: defaultdict(float))
    for r in SA:
        if r[1]==p and r[23] and r[8] in ('Nuoto','Ciclismo','Corsa a piedi'):
            mes[r[3]][r[8]]+=r[23]
    for m,d in mes.items():
        t=sum(d.values())
        print(f"    {m[:34]:34s} " + "  ".join(f"{k[:5]} {v/t*100:4.1f}%" for k,v in sorted(d.items())))

print("\n"+"="*78); print("E. BRICK / COMBINATI - dove il libro li segnala")
br=Counter()
for r in SA:
    if r[39] and 'combinare assieme' in str(r[39]): br[(r[1],r[7])]+=1
print("  ", dict(br) if br else "NESSUNA indicazione di combinato trovata")
print("  piani che li prevedono:", sorted({k[0] for k in br}) or "nessuno")

print("\n"+"="*78); print("F. DISTANZA FRA SEDUTE INTENSE (Z4-Z5)")
for p in PIANI:
    hard=defaultdict(set)
    for r in SA:
        if r[1]==p and r[30] in ('Z4','Z5') and r[23] and r[23]>3:
            hard[r[4]].add(r[7])
    ex=Counter()
    for s,gs in hard.items(): ex[tuple(sorted(gs, key=lambda g: GIORNI.index(g)))]+=1
    print(f"\n  {p}: giorni con lavoro Z4-Z5")
    for k,v in ex.most_common(3): print(f"    {v:2d} settimane -> {', '.join(k)}")

print("\n"+"="*78); print("G. PROGRESSIONE DENTRO IL MESOCICLO")
for p in PIANI:
    w=[r for r in OC if r[0]==p]
    print(f"\n  {p}: incrementi settimanali di carico (punti Edwards)")
    for i in range(1,len(w)):
        d=(w[i][5]/w[i-1][5]-1)*100
        print(f"    sett.{w[i-1][2]}->{w[i][2]}: {d:+6.1f}%", end="   ")
        if i%3==0: print()
    print()
